import json
from pathlib import Path

import openpyxl


def load_sidings(ws):
    sidings = []
    # Expect header row: Subdivision, Siding, Start MP, End MP, Length (ft)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        subdivision, name, start_mp, end_mp, length_ft = row
        sidings.append(
            {
                "subdivision": str(subdivision).strip() if subdivision else "",
                "name": str(name).strip() if name else "",
                "start_mp": float(start_mp) if start_mp is not None else None,
                "end_mp": float(end_mp) if end_mp is not None else None,
                "length_ft": int(length_ft) if length_ft is not None else None,
            }
        )
    return sidings


def main():
    # Master file location in your repo:
    xlsx_path = Path("data/public/RAILCORE_Master_Sidings.xlsx")
    if not xlsx_path.exists():
        raise SystemExit(f"Missing file: {xlsx_path}")

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    if "Sidings" not in wb.sheetnames:
        raise SystemExit("Workbook must contain a 'Sidings' sheet.")

    sidings_ws = wb["Sidings"]
    sidings = load_sidings(sidings_ws)

    out_dir = Path("data/json")
    out_dir.mkdir(parents=True, exist_ok=True)

    sidings_path = out_dir / "sidings.json"
    sidings_path.write_text(json.dumps({"sidings": sidings}, indent=2), encoding="utf-8")
    print(f"Wrote {sidings_path} with {len(sidings)} sidings.")


if __name__ == "__main__":
    main()
